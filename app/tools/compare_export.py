"""Structural comparison of an exported workbook vs the original Template sheet.

Usage:  python -m tools.compare_export <exported.xlsx> [sheet_name]
Prints a report and exits non-zero on structural differences outside the
intentionally regenerated body rows (7–28) and the data cells we fill.
"""
from __future__ import annotations

import sys

import openpyxl

from hrm import config

BODY_ROWS = range(7, 29)


def _cf(ws):
    out = []
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            out.append((str(cf.sqref), r.type, tuple(r.formula or ())))
    return sorted(out)


def _dv(ws):
    return sorted((str(d.sqref), d.type, d.formula1) for d in ws.data_validations.dataValidation)


def _style_sig(c):
    return (c.font.name, c.font.sz, c.font.b, c.font.color.rgb if c.font.color else None,
            c.fill.fgColor.rgb, c.number_format, c.alignment.horizontal,
            tuple(getattr(c.border, s).style for s in ("left", "right", "top", "bottom")))


def compare(export_path: str, sheet: str | None = None) -> list[str]:
    t = openpyxl.load_workbook(config.TEMPLATE_PATH)[config.TEMPLATE_SHEET]
    wb = openpyxl.load_workbook(export_path)
    e = wb[sheet] if sheet else wb.worksheets[0]
    problems, notes = [], []

    def chk(label, a, b):
        (notes if a == b else problems).append(f"{'OK ' if a == b else 'DIFF'} {label}: template={a!r} export={b!r}"
                                               if a != b else f"OK  {label}")

    chk("dimensions", t.dimensions, e.dimensions)
    chk("print_area", t.print_area.replace(f"'{t.title}'", "S"), e.print_area.replace(f"'{e.title}'", "S"))
    chk("orientation", t.page_setup.orientation, e.page_setup.orientation)
    chk("paperSize", t.page_setup.paperSize, e.page_setup.paperSize)
    chk("scale", t.page_setup.scale, e.page_setup.scale)
    mg = lambda w: tuple(round(getattr(w.page_margins, k), 6) for k in ("left", "right", "top", "bottom"))
    chk("margins", mg(t), mg(e))
    chk("freeze_panes", t.freeze_panes, e.freeze_panes)
    chk("column widths", {k: v.width for k, v in t.column_dimensions.items()},
        {k: v.width for k, v in e.column_dimensions.items()})
    chk("row heights", {k: v.height for k, v in t.row_dimensions.items()},
        {k: v.height for k, v in e.row_dimensions.items()})
    chk("conditional formatting", _cf(t), _cf(e))
    chk("data validation", _dv(t), _dv(e))
    tm = {str(m) for m in t.merged_cells.ranges if not (m.min_row >= 7 and m.max_row <= 28 and m.max_col <= 9)}
    em = {str(m) for m in e.merged_cells.ranges if not (m.min_row >= 7 and m.max_row <= 28 and m.max_col <= 9)}
    chk("merges outside body A7:I28", sorted(tm), sorted(em))

    # formulas & styles outside the body: must be identical
    fdiff, sdiff = [], []
    for row in t.iter_rows(min_row=1, max_row=46, max_col=48):
        for c in row:
            if c.row in BODY_ROWS:
                continue
            if 33 <= c.row <= 44 and (c.column == 8 or 10 <= c.column <= 43):
                continue   # shift data cells: static roster fills are reset on purpose
            ec = e[c.coordinate]
            if isinstance(c.value, str) and c.value.startswith("=") and c.value != ec.value:
                fdiff.append(c.coordinate)
            if _style_sig(c) != _style_sig(ec):
                sdiff.append(c.coordinate)
    chk("formulas outside body", [], fdiff)
    chk("cell styles outside body", [], sdiff)
    # body: style signature per column must come from the template's own body styles
    body_styles = {(c.column, _style_sig(c)[:3]) for r in t.iter_rows(min_row=7, max_row=28, max_col=45) for c in r}
    foreign = [c.coordinate for r in e.iter_rows(min_row=7, max_row=28, max_col=45) for c in r
               if (c.column, _style_sig(c)[:3]) not in body_styles]
    chk("body fonts drawn from template body", [], foreign[:20])
    codes = [v for row in e.iter_rows(min_row=33, max_row=44, min_col=10, max_col=43, values_only=True) for v in row]
    chk("shift day codes left blank (roster not imported)", True, all(v is None for v in codes))
    return problems + notes


if __name__ == "__main__":
    res = compare(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    for line in res:
        print(line)
    sys.exit(1 if any(l.startswith("DIFF") for l in res) else 0)
