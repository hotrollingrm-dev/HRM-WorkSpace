import pytest
import re
from pathlib import Path

import openpyxl

from conftest import needs_company_data, import_schedule, items
from hrm import config


@needs_company_data
def test_import_preview_against_real_file(client):
    b = import_schedule(client)
    rows = b["rows"]
    sheets = {s["sheet"]: s for s in b["summary"]["sheets"]}
    assert sheets["HOT Plan Apr"]["status"] == "skipped" and sheets["Master plan"]["status"] == "skipped"
    assert sheets["Oct"]["profile"] == "P2" and sheets["May"]["profile"] == "P1"
    assert any("August" in w for w in sheets["Jul"]["warnings"])          # Jul sheet titled August
    rej = [r for r in rows if r["action"] == "reject"]
    assert rej and all(r["sheet"] == "Jul" and r["cell_ref"].endswith("64") for r in rej)  # B64 'day 30' Fri
    assert all(r["date"] for r in rows)
    # multi-line cell → two events with their own times
    oct2 = [r for r in rows if r["sheet"] == "Oct" and r["cell_ref"] == "I6"]
    assert [(r["start_at"][11:16], r["end_at"][11:16]) for r in oct2] == [("13:00", "16:00"), ("19:40", "21:30")]
    # untimed line → unscheduled, never guessed
    un = [r for r in rows if not r["start_at"]]
    assert un and all(any(i["code"] == "unscheduled" for i in r["issues"]) for r in un)
    # numeric MBO flagged
    assert any(i["code"] == "mbo_numeric" for r in rows for i in r["issues"])
    # footer legend rows are not events
    assert not any("กิจกรรม MBO และอื่นๆ" == r["title"] for r in rows)


@needs_company_data
def test_import_commit_creates_draft_only_and_no_completion(client):
    b = import_schedule(client)
    r = client.post(f"/api/imports/{b['id']}/commit").json()
    assert r["added"] == sum(1 for x in b["rows"] if x["action"] == "add")
    v = client.get("/api/plans/2026/10").json()
    imp = [e for e in v["entries"] if e["source_type"] == "import"]
    assert imp and all(e["progress"]["status"] == "planned" for e in imp)
    assert all(re.fullmatch(r"Oct![A-Z]+\d+#\d+", e["source_ref"]) for e in imp)
    assert client.get("/api/today?d=2026-10-02").json()["entries"] == []     # not published yet


@needs_company_data
def test_reimport_skips_duplicates_and_respects_manual_edits(client):
    b = import_schedule(client)
    client.post(f"/api/imports/{b['id']}/commit")
    v = client.get("/api/plans/2026/10").json()
    target = next(e for e in v["entries"] if e["source_ref"] == "Oct!I6#1")
    client.patch(f"/api/entries/{target['uid']}", json={"start_time": "13:30", "end_time": "16:00"})
    b2 = import_schedule(client)
    acts = {r["action"] for r in b2["rows"]}
    assert "add" not in acts                                       # nothing duplicated
    edited = next(r for r in b2["rows"] if r["sheet"] == "Oct" and r["cell_ref"] == "I6" and r["line_no"] == 1)
    assert edited["action"] == "conflict"
    assert client.patch(f"/api/import-rows/{edited['id']}", json={"action": "update"}).status_code == 422
    client.post(f"/api/imports/{b2['id']}/commit")
    after = next(e for e in client.get("/api/plans/2026/10").json()["entries"] if e["uid"] == target["uid"])
    assert after["start_at"].endswith("13:30:00+07:00")            # manual edit preserved
    assert client.get(f"/api/imports/{b2['id']}").json()["summary"]["previous_commit_same_file"] == b["id"]


@needs_company_data
def test_rejected_row_fixed_by_user(client):
    b = import_schedule(client)
    rej = next(r for r in b["rows"] if r["action"] == "reject")
    r = client.patch(f"/api/import-rows/{rej['id']}", json={"date": "2026-07-31"}).json()
    assert r["action"] == "add" and r["date"] == "2026-07-31"


@needs_company_data
def test_revert_keeps_manually_edited(client):
    b = import_schedule(client)
    client.post(f"/api/imports/{b['id']}/commit")
    e = next(x for x in client.get("/api/plans/2026/10").json()["entries"] if x["source_type"] == "import")
    client.patch(f"/api/entries/{e['uid']}", json={"notes": "แก้แล้ว"})
    r = client.post(f"/api/imports/{b['id']}/revert").json()
    assert r["kept_manually_edited"] == [e["uid"]] and r["removed"] > 100
    left = [x for x in client.get("/api/plans/2026/10").json()["entries"] if x["source_type"] == "import"]
    assert [x["uid"] for x in left] == [e["uid"]]


def _export(client, **kw):
    r = client.post("/api/exports", json={"year": 2026, "month": 10, **kw})
    assert r.status_code == 200, r.text
    f = client.get(r.json()["url"])
    assert f.status_code == 200 and f.content[:2] == b"PK"
    return Path(config.EXPORT_DIR) / r.json()["file"]


@needs_company_data
def test_export_requires_publish_or_draft(client):
    w = items(client)[0]
    client.post("/api/plans/2026/10/entries", json={"dates": ["2026-10-01"], "work_item_id": w["id"]})
    assert client.post("/api/exports", json={"year": 2026, "month": 10}).status_code == 422
    assert _export(client, source="draft").name.endswith("_DRAFT.xlsx")


@needs_company_data
def test_export_matches_template_structure_and_marks(client):
    from tools.compare_export import compare
    its = items(client)
    for w in its:
        client.post("/api/plans/2026/10/rows", json={"work_item_ids": [w["id"]]})
    a, pm = its[0], next(w for w in its if w["source_ref"] == "Template 10!D25")
    uid = client.post("/api/plans/2026/10/entries", json={"dates": ["2026-10-01"], "work_item_id": a["id"]}).json()["created"][0]["uid"]
    client.post("/api/plans/2026/10/entries", json={"dates": ["2026-10-24"], "work_item_id": pm["id"]})
    b = import_schedule(client)
    client.post(f"/api/imports/{b['id']}/commit")
    client.post("/api/plans/2026/10/publish", json={})
    client.post(f"/api/progress/{uid}", json={"status": "completed"})
    path = _export(client)
    report = compare(str(path))
    assert not [l for l in report if l.startswith("DIFF")], report
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["10", "Daily Detail"]
    ws = wb["10"]
    assert ws["D2"].value.date().isoformat() == "2026-10-01"
    assert ws["M7"].value == "✔"                      # day 1 → column M, completed
    assert ws["AJ25"].value == "✔"                    # day 24 → column AJ, planned (✔ = แผน, user-confirmed)
    assert ws["AR25"].value == 1
    assert ws["AS7"].value == '=COUNTIF($M7:$AQ7,"✔")'
    assert ws["B7"].value == "ทำความสะอาด\nCleaning" and ws["D25"].value == pm["name_th"]
    merges = {str(m) for m in ws.merged_cells.ranges}
    t = openpyxl.load_workbook(config.TEMPLATE_PATH)["10"]
    assert {str(m) for m in t.merged_cells.ranges} == merges   # same item set → identical merges
    emps = client.get("/api/employees").json()
    by_team = sorted(emps, key=lambda e: (e["team"], e["sort_order"], e["id"]))
    assert [ws.cell(r, 3).value for r in range(33, 45)] == [e["name"] for e in by_team]
    assert [ws.cell(r, 8).value for r in range(33, 45)] == [e["team"] for e in by_team]
    assert all(ws.cell(r, c).value is None for r in range(33, 45) for c in range(10, 44))
    assert ws["AU22"].value.date().isoformat() == "2026-10-13"
    dd = wb["Daily Detail"]
    titles = [dd.cell(r, 7).value for r in range(5, dd.max_row + 1)]
    oct_i6 = next(r["title"] for r in b["rows"] if r["sheet"] == "Oct" and r["cell_ref"] == "I6" and r["line_no"] == 1)
    assert oct_i6 in titles
    assert "REV.01" in ws["F2"].value


@needs_company_data
def test_export_overflow_goes_to_detail_sheet(client):
    c = client.post("/api/categories", json={"name_th": "Extra"}).json()
    s = client.post("/api/subcategories", json={"category_id": c["id"], "name_th": "X"}).json()
    ids = [w["id"] for w in items(client)]
    for i in range(3):
        ids.append(client.post("/api/items", json={"subcategory_id": s["id"], "name_th": f"Extra {i}"}).json()["id"])
    client.post("/api/plans/2026/10/rows", json={"work_item_ids": ids})
    extra = ids[-1]
    client.post("/api/plans/2026/10/entries", json={"dates": ["2026-10-09"], "work_item_id": extra})
    path = _export(client, source="draft")
    wb = openpyxl.load_workbook(path)
    assert "Extra 2" in [wb["Daily Detail"].cell(r, 7).value for r in range(5, 40)]
    assert wb["10"].max_row == 46


def test_no_ot_label_in_ui():
    web = Path(__file__).resolve().parent.parent / "web"
    for f in list((web / "js").glob("*.js")) + list((web / "css").glob("*.css")) + list(web.glob("*.html")):
        assert not re.search(r"\bOT\b", f.read_text(encoding="utf-8")), f.name


def test_every_ui_label_bilingual():
    src = (Path(__file__).resolve().parent.parent / "web" / "js" / "ui.js").read_text(encoding="utf-8")
    pairs = re.findall(r'\w+: \["([^"]+)", "([^"]*)"\]', src.split("export const T = {")[1].split("};")[0])
    assert len(pairs) > 30 and all(th and jp for th, jp in pairs)


@needs_company_data
def test_static_pages_served(client):
    for p in ("/", "/m.html", "/manifest.webmanifest", "/sw.js", "/icons/icon-192.png"):
        assert client.get(p).status_code == 200, p


@needs_company_data
def test_export_reflects_team_change(client):
    w = items(client)[0]
    client.post("/api/plans/2026/10/entries", json={"dates": ["2026-10-01"], "work_item_id": w["id"]})
    e = client.get("/api/employees").json()[0]                    # first of team A
    client.patch(f"/api/employees/{e['id']}", json={"team": "C"})
    ws = openpyxl.load_workbook(_export(client, source="draft"))["10"]
    rows = {ws.cell(r, 3).value: ws.cell(r, 8).value for r in range(33, 45)}
    assert rows[e["name"]] == "C"
    assert [ws.cell(r, 8).value for r in range(33, 45)] == ["A"] * 3 + ["B"] * 4 + ["C"] * 5   # grouped by team
