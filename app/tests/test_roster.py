import calendar

import openpyxl

from conftest import needs_company_data, items
from hrm.services.roster import _Src, _source_cells, run_checks


@needs_company_data
def test_transcription_matches_pdf_checksums():
    c = run_checks(_source_cells())
    assert c["wd_rows_checked"] == 45 and c["wd_mismatches"] == []
    assert all(v["match"] for v in c["summary_2026"].values())          # D/N/H/S per group = PDF summary box
    assert all(x["explained_by_substitution"] for x in c["coverage_exceptions"])   # uneven days = green cells only
    assert c["ok"] is True
    assert c["shutdown_days_2026"] == 11                                # user: follow the table (not the "10 days" note)
    SRC = _Src()
    for ym, g in SRC.CODES.items():
        y, m = map(int, ym.split("-"))
        assert all(len(g[t]) == calendar.monthrange(y, m)[1] and set(g[t]) <= set("DNHS") for t in "ABC")


def _normal_day(cells):
    """A date with exactly one D, one N and one H team (expectations come from company_data, not code)."""
    return next(d for d, v in sorted(cells.items())
                if d >= "2026-10-01" and sorted(x[0] for x in v.values()) == ["D", "H", "N"] and not any(x[1] for x in v.values()))


@needs_company_data
def test_fills_detected_from_pixels():
    cells = _source_cells()
    kinds = [(v[0], v[1]) for day in cells.values() for v in day.values() if v[1]]
    assert {k for _, k in kinds} == {"special", "substitution"}
    assert all(code in "DNS" for code, k in kinds if k == "special")     # yellow = working holiday or shutdown
    assert any(code == "S" for code, k in kinds if k == "special")


@needs_company_data
def test_roster_needs_confirmation_then_drives_today_and_export(client):
    cells = _source_cells()
    day = _normal_day(cells)
    team_of = {code: t for t, (code, _) in cells[day].items()}
    t = client.get(f"/api/today?d={day}").json()
    assert t["shift"]["status"] == "no_calendar"
    pv = client.post("/api/roster/preview").json()
    assert pv["status"] == "preview" and pv["checks"]["ok"]
    assert client.get(f"/api/today?d={day}").json()["shift"]["status"] == "no_calendar"   # not before commit
    r = client.post(f"/api/roster/{pv['id']}/commit").json()
    assert r["added"] == 455 * 3
    t = client.get(f"/api/today?d={day}").json()["shift"]
    d, n = t["shifts"]
    assert (d["code"], d["team"], n["code"], n["team"]) == ("D", team_of["D"], "N", team_of["N"])
    from datetime import date as _d, timedelta as _td
    nxt = (_d.fromisoformat(day) + _td(days=1)).isoformat()
    assert d["start"] == f"{day}T08:00:00+07:00" and d["end"] == f"{day}T20:15:00+07:00"
    assert n["start"] == f"{day}T20:00:00+07:00" and n["end"] == f"{nxt}T08:15:00+07:00"
    team_d = [e["name"] for e in client.get("/api/employees").json() if e["team"] == team_of["D"]]
    assert [m["name"] for m in d["members"]] == team_d and team_d
    assert t["off"] == [{"team": team_of["H"], "code": "H"}]
    assert client.get("/api/today?d=2027-01-10").json()["shift"]["provisional"] is True
    # export fills D/N for staff rows by team
    w = items(client)[0]
    client.post("/api/plans/2026/10/entries", json={"dates": ["2026-10-01"], "work_item_id": w["id"]})
    f = client.post("/api/exports", json={"year": 2026, "month": 10, "source": "draft"}).json()["file"]
    from hrm import config
    ws = openpyxl.load_workbook(config.EXPORT_DIR / f)["10"]
    tmpl = openpyxl.load_workbook(config.TEMPLATE_PATH)["10"]
    # The template's own roster for Oct (rows 33–44, J:AQ) must equal what we generate from the PDF
    same = [(r, c) for r in range(33, 45) for c in range(10, 44) if ws.cell(r, c).value != tmpl.cell(r, c).value]
    assert same == []
    assert ws["AR33"].value == '=COUNTIF($M33:$AQ33,"D")'


@needs_company_data
def test_manual_day_edit_survives_reimport(client):
    pv = client.post("/api/roster/preview").json()
    client.post(f"/api/roster/{pv['id']}/commit")
    assert client.patch("/api/roster/2026-10-20/A", json={"code": "D"}).json()["edited"] is True
    pv2 = client.post("/api/roster/preview").json()
    r = client.post(f"/api/roster/{pv2['id']}/commit").json()
    assert r["added"] == 0 and r["kept_manually_edited"] == 1
    m = client.get("/api/roster?year=2026&month=10").json()
    assert m["days"]["2026-10-20"]["A"] == {"code": "D", "fill": None, "edited": True}
    assert client.patch("/api/roster/2026-10-20/A", json={"code": "X"}).status_code == 422


@needs_company_data
def test_official_times_confirmed(client):
    s = client.get("/api/settings").json()["official_roster_times"]
    assert s["confirmed"] is True and s["D"]["end"] == "20:15" and s["N"]["end"] == "08:15"


@needs_company_data
def test_special_fill_means_public_holiday_work(client):
    pv = client.post("/api/roster/preview").json()
    client.post(f"/api/roster/{pv['id']}/commit")
    cells = _source_cells()
    day = next(d for d, v in sorted(cells.items())
               if sorted(c for c, f in v.values() if f == "special") == ["D", "N"])   # a public holiday worked by both shifts
    sh = client.get(f"/api/today?d={day}").json()["shift"]
    assert {x["code"]: x["fill"] for x in sh["shifts"]} == {"D": "special", "N": "special"}
    from hrm.services.roster import FILL_TH
    assert FILL_TH["special"].startswith("ทำงานวันหยุดนักขัตฤกษ์")
